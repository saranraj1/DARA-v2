"""Generate DARA Admin UI specification document as a .docx file."""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Title
doc.add_heading("DARA Admin UI — Feature & Page Specification", 0)
doc.add_paragraph(
    "This document describes every page, tab, section, and interactive element "
    "of the DARA Admin UI. Use this as a prompt for generating the interface."
)

# ── OVERVIEW ──────────────────────────────────────────────────
doc.add_heading("Overview", 1)
doc.add_paragraph(
    "DARA (Distributed Architecture Root-cause Agent) is an autonomous AI debugging system. "
    "The Admin UI is a single-page web application with a persistent left sidebar for navigation "
    "and a main content area that changes per page. All data comes from a live REST API backend. "
    "The UI must refresh data automatically every 30 seconds on every page."
)

doc.add_heading("Global Elements", 2)
items = [
    "Left sidebar with navigation links to all pages.",
    "App name and version number at the top of the sidebar.",
    "A live connection indicator (dot + label) showing whether the backend API is reachable.",
    "A Refresh button on every page header to manually reload data.",
    "Last-updated timestamp displayed on every page header.",
    "A link to the API documentation (Swagger UI) at the bottom of the sidebar.",
]
for item in items:
    doc.add_paragraph(item, style="List Bullet")

# ── PAGE 1: DASHBOARD ─────────────────────────────────────────
doc.add_heading("Page 1 — Dashboard", 1)
doc.add_paragraph(
    "The Dashboard is the landing page. It gives a high-level summary of the entire DARA pipeline."
)

doc.add_heading("KPI Cards (top row)", 2)
kpis = [
    "Total Errors — total number of errors ingested into the system.",
    "Fixed — number of errors that have an accepted fix.",
    "Pending — number of errors awaiting a fix or HITL review.",
    "Escalated — number of errors marked as escalated (no fix found).",
    "Avg Confidence — average confidence score of all generated fixes (shown as a percentage).",
    "API Health — a visual health bar showing whether the backend is healthy.",
]
for k in kpis:
    doc.add_paragraph(k, style="List Bullet")

doc.add_heading("Charts Section", 2)
charts = [
    "Pipeline Throughput chart — a time-series area chart showing errors ingested vs fixes generated over the last hour in 5-minute buckets.",
    "Errors by Class chart — a bar chart showing the count of errors grouped by error class (e.g. null_reference, timeout, syntax_error).",
    "Recent Outcomes — a list of the 5 most recent pipeline outcomes (error message, service, fix outcome, timestamp).",
]
for c in charts:
    doc.add_paragraph(c, style="List Bullet")

# ── PAGE 2: ERRORS ────────────────────────────────────────────
doc.add_heading("Page 2 — Errors", 1)
doc.add_paragraph(
    "The Errors page shows all errors ingested into the system in a filterable, paginated table."
)

doc.add_heading("Filter Bar", 2)
filters = [
    "Status dropdown — filter by: All, Pending, Analyzing, Fixed, Escalated, Ignored.",
    "Severity dropdown — filter by: All, Critical, High, Medium, Low.",
    "Service name text input — filter by service name (partial match).",
    "Error class text input — filter by error class.",
    "Date range picker — filter errors by ingestion date range.",
]
for f in filters:
    doc.add_paragraph(f, style="List Bullet")

doc.add_heading("Errors Table", 2)
cols = [
    "Error Class — the category of the error (e.g. null_reference).",
    "Message — the first 80 characters of the error message.",
    "Service — the name of the microservice that produced the error.",
    "Severity — badge showing critical / high / medium / low.",
    "Status — badge showing current pipeline state.",
    "Created At — human-readable relative timestamp (e.g. '2 hours ago').",
    "Actions — a button to view full error details.",
]
for c in cols:
    doc.add_paragraph(c, style="List Bullet")

doc.add_heading("Error Detail Panel (slide-in or modal)", 2)
doc.add_paragraph(
    "Clicking an error opens a detail panel with: full error message, full stack trace, "
    "service name, repository, commit SHA, pipeline run ID, list of related fixes, "
    "and a link to the originating GitHub Actions run."
)

doc.add_heading("Pagination", 2)
doc.add_paragraph("Page size selector (10, 20, 50) and previous/next page controls.")

# ── PAGE 3: FIXES ─────────────────────────────────────────────
doc.add_heading("Page 3 — Fixes (HITL Queue)", 1)
doc.add_paragraph(
    "The Fixes page is the Human-in-the-Loop approval interface. It lists all AI-generated fixes "
    "waiting for review, already approved, or rejected."
)

doc.add_heading("Filter Bar", 2)
doc.add_paragraph("Outcome filter dropdown: All, Pending, Accepted, Rejected, Auto-merged.")

doc.add_heading("Fixes Table", 2)
cols = [
    "Fix ID — shortened UUID.",
    "Error — the error class this fix targets.",
    "Strategy — the fix strategy used (e.g. llm_single_file, llm_multi_file).",
    "Confidence — a progress bar showing the AI confidence score (0–100%).",
    "Validation — pass / fail / pending badge from the sandbox test run.",
    "Outcome — current HITL outcome badge.",
    "PR Link — a clickable link to the GitHub Pull Request (if created).",
    "Files Changed — comma-separated list of files the fix modifies.",
    "Actions — Approve button and Reject button (only shown when outcome is pending).",
]
for c in cols:
    doc.add_paragraph(c, style="List Bullet")

doc.add_heading("Fix Detail Panel (slide-in or modal)", 2)
doc.add_paragraph(
    "Clicking a fix shows: the full generated diff (syntax-highlighted code diff view), "
    "the LLM's root cause explanation, the validation output (test results, linter output), "
    "the confidence score breakdown, and the HITL Approve / Reject buttons with an optional comment field."
)

# ── PAGE 4: PATTERNS ──────────────────────────────────────────
doc.add_heading("Page 4 — Pattern Library", 1)
doc.add_paragraph(
    "The Pattern Library shows recurring error patterns that DARA has learned from accepted fixes. "
    "Each pattern represents a solved class of problems."
)

doc.add_heading("Pattern Cards", 2)
fields = [
    "Pattern name / error class.",
    "Description — what this pattern detects.",
    "Match count — how many times this pattern has been matched.",
    "Success rate — percentage of fixes for this pattern that were accepted.",
    "Embedding similarity threshold — confidence score required for a match.",
    "Last matched — relative timestamp.",
    "Status badge — active / deprecated.",
    "Edit button — allows updating the pattern description or threshold.",
]
for f in fields:
    doc.add_paragraph(f, style="List Bullet")

doc.add_heading("Empty State", 2)
doc.add_paragraph(
    "When no patterns exist, show a message explaining that patterns are automatically created "
    "after fixes are accepted."
)

# ── PAGE 5: STRATEGY ──────────────────────────────────────────
doc.add_heading("Page 5 — Strategy Variants", 1)
doc.add_paragraph(
    "The Strategy page shows the A/B tested LLM prompt strategies DARA uses to generate fixes. "
    "Each error class can have multiple strategy variants being tested."
)

doc.add_heading("Strategy Table", 2)
cols = [
    "Error Class — the type of error this strategy targets.",
    "Variant Name — the name of this prompt strategy.",
    "Status badge — draft / testing / active / retired.",
    "Acceptance Rate — percentage of fixes accepted when using this strategy.",
    "Total Uses — how many times this strategy has been used.",
    "Success Rate — percentage shown as a progress bar.",
    "Generated By — 'llm' or 'human'.",
    "Created At — date the strategy was created.",
    "Actions — Promote to Active button (if testing), Retire button (if active).",
]
for c in cols:
    doc.add_paragraph(c, style="List Bullet")

doc.add_heading("Active Strategy Indicator", 2)
doc.add_paragraph(
    "A prominent banner at the top of the page showing the currently active strategy for each "
    "error class and its live success rate."
)

# ── PAGE 6: TOPOLOGY ──────────────────────────────────────────
doc.add_heading("Page 6 — Service Topology", 1)
doc.add_paragraph(
    "The Topology page visualises the live service call graph built from ingested "
    "OpenTelemetry traces. It shows which microservices call each other and highlights "
    "services with high error rates."
)

doc.add_heading("Service Graph (main panel)", 2)
doc.add_paragraph(
    "An interactive node-link graph where each node is a service and each directed edge "
    "is a call relationship. Nodes are clickable. Edges show call count and error rate on hover. "
    "Services with error rate above 5% have a visual alert indicator on their node."
)

doc.add_heading("Graph Controls", 2)
controls = [
    "Zoom in / zoom out buttons.",
    "Reset view button (re-centres the graph).",
    "Filter input — hide services not matching the typed name.",
]
for c in controls:
    doc.add_paragraph(c, style="List Bullet")

doc.add_heading("Service Detail Panel (on node click)", 2)
fields = [
    "Service name.",
    "Total call count (incoming + outgoing).",
    "Error count and error rate.",
    "Average latency (ms).",
    "Last seen timestamp.",
    "List of all services this service calls (outgoing edges).",
    "List of all services that call this service (incoming edges).",
    "Link to filter the Errors page by this service.",
]
for f in fields:
    doc.add_paragraph(f, style="List Bullet")

doc.add_heading("Recent Traces Table (below the graph)", 2)
cols = [
    "Trace ID — shortened, clickable.",
    "Span Count — number of spans in the trace.",
    "Services Involved — comma-separated list.",
    "Has Error — yes/no badge.",
    "Started At — relative timestamp.",
    "Duration — total trace duration in ms.",
]
for c in cols:
    doc.add_paragraph(c, style="List Bullet")

doc.add_heading("Empty State", 2)
doc.add_paragraph(
    "When no traces have been ingested, show a message explaining how to connect "
    "an OpenTelemetry-instrumented service to DARA."
)

# ── PAGE 7: AUDIT LOG ─────────────────────────────────────────
doc.add_heading("Page 7 — Audit Log", 1)
doc.add_paragraph(
    "The Audit Log records every HITL decision (approve / reject) made on fixes, "
    "providing a full history of human interventions."
)

doc.add_heading("Audit Table", 2)
cols = [
    "Timestamp — exact date and time of the decision.",
    "Action — Approved or Rejected badge.",
    "Fix ID — shortened UUID, links to the fix detail.",
    "Error Class — the type of error.",
    "Reviewer — who made the decision (username or 'system' for auto-merge).",
    "Comment — optional comment left by the reviewer.",
    "Confidence At Decision — the AI confidence score at the time of the HITL decision.",
]
for c in cols:
    doc.add_paragraph(c, style="List Bullet")

doc.add_heading("Filter Bar", 2)
filters = [
    "Action filter: All, Approved, Rejected.",
    "Date range picker.",
    "Reviewer input (text search).",
]
for f in filters:
    doc.add_paragraph(f, style="List Bullet")

doc.add_heading("Export Button", 2)
doc.add_paragraph("A button to export the filtered audit log as a CSV file.")

# ── API ENDPOINTS USED ─────────────────────────────────────────
doc.add_heading("Backend API Endpoints (for reference)", 1)
endpoints = [
    "GET  /health                        — backend health check",
    "GET  /health/ready                  — full readiness check",
    "GET  /api/v1/admin/stats            — dashboard KPI numbers",
    "GET  /api/v1/admin/errors           — paginated error list (supports filters)",
    "GET  /api/v1/admin/fixes            — paginated fix list (supports filters)",
    "GET  /api/v1/admin/patterns         — pattern library list",
    "GET  /api/v1/admin/audit            — audit log list",
    "POST /api/v1/fixes/{id}/approve     — approve a fix",
    "POST /api/v1/fixes/{id}/reject      — reject a fix",
    "GET  /api/v1/topology               — service topology graph data",
    "GET  /api/v1/traces                 — list of recent traces",
    "GET  /api/v1/traces/{trace_id}      — full trace tree",
]
for e in endpoints:
    doc.add_paragraph(e, style="List Bullet")

doc.add_heading("Authentication", 2)
doc.add_paragraph(
    "All /api/v1/admin/* endpoints require the header: X-Admin-Token: <secret>. "
    "The public base URL during development is http://localhost:8000. "
    "The production public URL is https://tremor-laptop-prepay.ngrok-free.dev."
)

# ── INTERACTIONS ──────────────────────────────────────────────
doc.add_heading("Key Interactions and Behaviours", 1)
interactions = [
    "Approving a fix: click Approve → confirmation dialog → POST to approve endpoint → row updates to 'Accepted' without page reload.",
    "Rejecting a fix: click Reject → optional comment dialog → POST to reject endpoint → row updates to 'Rejected' without page reload.",
    "Clicking a service node in the topology graph opens the service detail panel inline.",
    "Clicking a trace ID in the traces table expands a full span tree view inline.",
    "All tables support sorting by clicking column headers.",
    "All tables show a loading skeleton while data is being fetched.",
    "All tables show a clear empty state message when there is no data.",
    "API errors show a dismissable error banner at the top of the page with the HTTP status code.",
    "The app does not require login — it uses a shared admin token configured at build time.",
    "The sidebar navigation highlights the currently active page.",
]
for i in interactions:
    doc.add_paragraph(i, style="List Bullet")

# Save
output_path = r"c:\Production level projects\ADAA\DARA_Admin_UI_Specification.docx"
doc.save(output_path)
print(f"Saved: {output_path}")
